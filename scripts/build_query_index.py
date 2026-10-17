#!/usr/bin/env python3
from pathlib import Path
import argparse, hashlib, json, re, sqlite3, sys
import yaml
from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from afc_card_runtime.common import read_jsonl, normalize_text, tokenize, TOKENIZER_VERSION
from afc_card_runtime.scope import load_business_context_map as load_context_rows

FIELD_SEARCH=['名稱','中文','越文','中文補充','說明','規則','原始商品資訊','商品','商品分類','商品代號','功能代碼','System','資料來源','Related Field','來源','Supporting Source']
CARD_SEARCH=['Title','Content','Original Evidence','中文補充','Scope','Product Context','Context','Card Type','Concept Key','Semantic Domain Keys','Semantic Domain Labels','Source','Source Location','Evidence Type','Notes']
EVIDENCE_SEARCH=['Original Text','Source Language','Original Name','Original Description','Original Message','Original Component Type','Original Code','Original Value','Original Label','Original Placeholder','Original Required','Original Max Length','Original Min Length','Original Format','Original Validation Message','Original Error Message','Original Option Values','Original Function Code','Structured Facts','Evidence Kind','Source','Source Location']
SHARED_SCOPES={'Multiple','多商品','Shared','shared'}


def sha256_file(path):
    p=Path(path) if path else None
    if not p or not p.exists(): return None
    h=hashlib.sha256()
    with p.open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024), b''): h.update(chunk)
    return h.hexdigest()


def search_tokens(text):
    """Build the FTS5 search_text for one record.

    The raw text is deliberately NOT included: FTS5 `unicode61` collapses an
    unbroken CJK run into a single token, so storing the raw string only added a
    term that no short query could ever match. tokenize() now emits overlapping
    CJK n-grams, which is what makes Chinese retrieval work at all.
    """
    text=normalize_text(text).lower(); toks=tokenize(text)
    codes=re.findall(r'(?<![A-Za-z0-9_])[A-Za-z]{1,8}[0-9]{0,4}(?![A-Za-z0-9_])', text)
    vals=[]
    for x in [*toks,*codes]:
        x=normalize_text(x).replace('"',' ')
        if x and x not in vals: vals.append(x)
    return ' '.join(vals)


def load_field_xlsx(path):
    if not path: return []
    p=Path(path)
    if not p.exists(): return []
    wb=load_workbook(p,read_only=True,data_only=True)
    if 'Field_Data' not in wb.sheetnames: wb.close(); return []
    ws=wb['Field_Data']; it=ws.iter_rows(values_only=True)
    try: headers=[str(x) if x is not None else '' for x in next(it)]
    except StopIteration: wb.close(); return []
    rows=[{headers[i]:vals[i] for i in range(len(headers))} for vals in it]
    wb.close(); return rows


def item_text(row,fields):
    parts=[]
    for f in fields:
        v=row.get(f)
        if isinstance(v,list): v=' '.join(map(str,v))
        elif isinstance(v,dict): v=json.dumps(v,ensure_ascii=False,default=str)
        v=normalize_text(v)
        if v: parts.append(v)
    return ' | '.join(parts)


def scopes_from_card(row, attributed=None):
    if attributed is not None:
        return list(dict.fromkeys(normalize_text(x) for x in attributed if normalize_text(x)))
    v=row.get('Scopes')
    if v is None: v=row.get('Scope')
    if v is None: v=row.get('Product Context')
    vals=[]
    if isinstance(v,list): raw=v
    elif v is None: raw=[]
    else: raw=re.split(r'[,;，；]+', normalize_text(v))
    for x in raw:
        x=normalize_text(x)
        if x and x not in vals: vals.append(x)
    return vals


def load_business_context_map(path):
    if not path or not Path(path).exists(): return []
    rows=[]
    for item in load_context_rows(path):
        for alias in item['aliases']:
            for scope in item['scopes']:
                rows.append((item['context'],alias,alias.lower(),scope,item['description']))
    return rows


def db_connect(path):
    p=Path(path); p.parent.mkdir(parents=True,exist_ok=True)
    con=sqlite3.connect(p)
    con.execute('PRAGMA journal_mode=WAL'); con.execute('PRAGMA synchronous=NORMAL'); con.execute('PRAGMA temp_store=MEMORY')
    return con


def build(a):
    con=db_connect(a.output); cur=con.cursor()
    domain_terms={}
    if a.domains:
        for domain in read_jsonl(a.domains):
            key=normalize_text(domain.get('Domain Key'))
            if key:
                domain_terms[key]=' '.join([str(domain.get('Canonical Label') or ''), *[str(x) for x in domain.get('Aliases') or []]])
    cur.executescript('''
    DROP TABLE IF EXISTS meta; DROP TABLE IF EXISTS field_items; DROP TABLE IF EXISTS card_items; DROP TABLE IF EXISTS card_scopes;
    DROP TABLE IF EXISTS evidence_items; DROP TABLE IF EXISTS evidence_card;
    DROP TABLE IF EXISTS links; DROP TABLE IF EXISTS business_context_map; DROP TABLE IF EXISTS boxes; DROP TABLE IF EXISTS box_alias_map;
    DROP TABLE IF EXISTS box_edges; DROP TABLE IF EXISTS box_hierarchy; DROP TABLE IF EXISTS box_closure; DROP TABLE IF EXISTS card_box;
    DROP TABLE IF EXISTS field_fts; DROP TABLE IF EXISTS card_fts; DROP TABLE IF EXISTS evidence_fts;
    CREATE TABLE meta(key TEXT PRIMARY KEY,value TEXT);
    CREATE TABLE field_items(id TEXT PRIMARY KEY,product TEXT,product_category TEXT,candidate_group TEXT,name TEXT,search_text TEXT,payload TEXT);
    CREATE INDEX idx_field_product ON field_items(product); CREATE INDEX idx_field_group ON field_items(candidate_group);
    CREATE TABLE card_items(id TEXT PRIMARY KEY,primary_scope TEXT,card_type TEXT,title TEXT,concept_key TEXT,context TEXT,source TEXT,is_concept INTEGER DEFAULT 0,is_shared INTEGER DEFAULT 0,status TEXT,search_text TEXT,payload TEXT);
    CREATE INDEX idx_card_type ON card_items(card_type); CREATE INDEX idx_card_concept ON card_items(concept_key); CREATE INDEX idx_card_status ON card_items(status);
    CREATE TABLE card_scopes(card_id TEXT,scope TEXT,PRIMARY KEY(card_id,scope)); CREATE INDEX idx_card_scope ON card_scopes(scope);
    CREATE TABLE evidence_items(id TEXT PRIMARY KEY,evidence_mode TEXT,evidence_kind TEXT,source_language TEXT,source TEXT,source_location TEXT,status TEXT,search_text TEXT,payload TEXT);
    CREATE INDEX idx_evidence_kind ON evidence_items(evidence_kind); CREATE INDEX idx_evidence_language ON evidence_items(source_language); CREATE INDEX idx_evidence_status ON evidence_items(status);
    CREATE TABLE evidence_card(evidence_id TEXT,card_id TEXT,PRIMARY KEY(evidence_id,card_id)); CREATE INDEX idx_evidence_card_card ON evidence_card(card_id);
    CREATE TABLE business_context_map(context TEXT,alias TEXT,normalized_alias TEXT,scope TEXT,description TEXT,PRIMARY KEY(context,alias,scope));
    CREATE INDEX idx_business_context_alias ON business_context_map(normalized_alias); CREATE INDEX idx_business_context_scope ON business_context_map(scope);
    CREATE TABLE boxes(id TEXT PRIMARY KEY,box_type TEXT,name TEXT,canonical_key TEXT,description TEXT,status TEXT,scope_hints TEXT,payload TEXT);
    CREATE INDEX idx_box_type ON boxes(box_type); CREATE UNIQUE INDEX idx_box_key ON boxes(canonical_key);
    CREATE TABLE box_alias_map(box_id TEXT,alias TEXT,normalized_alias TEXT,PRIMARY KEY(box_id,alias)); CREATE INDEX idx_box_alias ON box_alias_map(normalized_alias);
    CREATE TABLE box_edges(id TEXT PRIMARY KEY,source_id TEXT,relation_type TEXT,target_id TEXT,link_basis TEXT,evidence TEXT,status TEXT,payload TEXT);
    CREATE INDEX idx_box_edge_source ON box_edges(source_id); CREATE INDEX idx_box_edge_target ON box_edges(target_id); CREATE INDEX idx_box_edge_type ON box_edges(relation_type);
    CREATE TABLE box_hierarchy(parent_id TEXT,child_id TEXT,source TEXT,PRIMARY KEY(parent_id,child_id));
    CREATE TABLE box_closure(ancestor_id TEXT,descendant_id TEXT,depth INTEGER,PRIMARY KEY(ancestor_id,descendant_id));
    CREATE INDEX idx_box_closure_ancestor ON box_closure(ancestor_id); CREATE INDEX idx_box_closure_desc ON box_closure(descendant_id);
    CREATE TABLE card_box(card_id TEXT,box_id TEXT,membership_type TEXT,basis TEXT,evidence TEXT,status TEXT,payload TEXT,PRIMARY KEY(card_id,box_id,membership_type));
    CREATE INDEX idx_card_box_box ON card_box(box_id); CREATE INDEX idx_card_box_card ON card_box(card_id);
    CREATE TABLE links(id TEXT PRIMARY KEY,source_id TEXT,target_id TEXT,link_type TEXT,link_basis TEXT,evidence TEXT,payload TEXT);
    CREATE INDEX idx_link_source ON links(source_id); CREATE INDEX idx_link_target ON links(target_id); CREATE INDEX idx_link_type ON links(link_type);
    CREATE VIRTUAL TABLE field_fts USING fts5(id UNINDEXED,search_text,tokenize='unicode61');
    CREATE VIRTUAL TABLE card_fts USING fts5(id UNINDEXED,search_text,tokenize='unicode61');
    CREATE VIRTUAL TABLE evidence_fts USING fts5(id UNINDEXED,search_text,tokenize='unicode61');
    ''')

    fields=load_field_xlsx(a.field_xlsx)
    cards=read_jsonl(a.cards) if a.cards and Path(a.cards).exists() else []
    scope_attributions=read_jsonl(a.scope_attributions) if a.scope_attributions and Path(a.scope_attributions).exists() else None
    evidence=read_jsonl(a.evidence) if a.evidence and Path(a.evidence).exists() else []
    links=read_jsonl(a.links) if a.links and Path(a.links).exists() else []
    boxes=read_jsonl(a.boxes) if a.boxes and Path(a.boxes).exists() else []
    memberships=read_jsonl(a.box_memberships) if a.box_memberships and Path(a.box_memberships).exists() else []
    box_links=read_jsonl(a.box_links) if a.box_links and Path(a.box_links).exists() else []
    ctx_rows=load_business_context_map(a.context_map)

    for r in fields:
        fid=normalize_text(r.get('Field ID'))
        if not fid: continue
        st=search_tokens(item_text(r,FIELD_SEARCH)); product=normalize_text(r.get('商品代號') or r.get('商品'))
        cur.execute('INSERT OR REPLACE INTO field_items VALUES(?,?,?,?,?,?,?)',(fid,product,normalize_text(r.get('商品分類')),normalize_text(r.get('Candidate Group')),normalize_text(r.get('名稱')),st,json.dumps(r,ensure_ascii=False,default=str)))
        cur.execute('INSERT INTO field_fts(id,search_text) VALUES(?,?)',(fid,st))

    attributed_scopes={}
    if scope_attributions is not None:
        for row in scope_attributions:
            if normalize_text(row.get('Status'))!='Confirmed': continue
            cid=normalize_text(row.get('Card ID')); scope=normalize_text(row.get('Scope ID'))
            if cid and scope: attributed_scopes.setdefault(cid,[]).append(scope)

    for r in cards:
        cid=normalize_text(r.get('Card ID'))
        if not cid: continue
        scopes=scopes_from_card(r, attributed_scopes.get(cid,[]) if scope_attributions is not None else None); primary=scopes[0] if scopes else ''
        ctype=normalize_text(r.get('Card Type'))
        tag_terms=' '.join(domain_terms.get(normalize_text(key),'') for key in (r.get('Semantic Domain Keys') or []))
        st=search_tokens(item_text(r,CARD_SEARCH)+' '+tag_terms)
        is_shared=1 if any(s in SHARED_SCOPES for s in scopes) else 0
        cur.execute('INSERT OR REPLACE INTO card_items VALUES(?,?,?,?,?,?,?,?,?,?,?,?)',(cid,primary,ctype,normalize_text(r.get('Title')),normalize_text(r.get('Concept Key')),normalize_text(r.get('Context')),normalize_text(r.get('Source')),1 if ctype=='Concept' else 0,is_shared,normalize_text(r.get('Status')) or 'Active',st,json.dumps(r,ensure_ascii=False,default=str)))
        for s in scopes: cur.execute('INSERT OR IGNORE INTO card_scopes VALUES(?,?)',(cid,s))
        cur.execute('INSERT INTO card_fts(id,search_text) VALUES(?,?)',(cid,st))

    # Source-preserving Evidence Records are indexed separately from normalized Card claims.
    evidence_links=set()
    for r in evidence:
        eid=normalize_text(r.get('Evidence ID'))
        if not eid: continue
        st=search_tokens(item_text(r,EVIDENCE_SEARCH))
        cur.execute('INSERT OR REPLACE INTO evidence_items VALUES(?,?,?,?,?,?,?,?,?)',(eid,normalize_text(r.get('Evidence Mode')),normalize_text(r.get('Evidence Kind')),normalize_text(r.get('Source Language')),normalize_text(r.get('Source')),normalize_text(r.get('Source Location')),normalize_text(r.get('Status')) or 'Active',st,json.dumps(r,ensure_ascii=False,default=str)))
        cur.execute('INSERT INTO evidence_fts(id,search_text) VALUES(?,?)',(eid,st))
        for cid in r.get('Card IDs') or []:
            cid=normalize_text(cid)
            if cid: evidence_links.add((eid,cid))
    # Also accept the reverse Card -> Evidence Record IDs declaration.
    for r in cards:
        cid=normalize_text(r.get('Card ID'))
        for eid in r.get('Evidence Record IDs') or []:
            eid=normalize_text(eid)
            if eid and cid: evidence_links.add((eid,cid))
    for eid,cid in evidence_links:
        cur.execute('INSERT OR IGNORE INTO evidence_card VALUES(?,?)',(eid,cid))

    for row in ctx_rows: cur.execute('INSERT OR REPLACE INTO business_context_map VALUES(?,?,?,?,?)',row)

    known_box_ids=set(); legacy_hierarchy=[]; legacy_members=[]
    for b in boxes:
        bid=normalize_text(b.get('Box ID'))
        if not bid: continue
        known_box_ids.add(bid); name=normalize_text(b.get('Name')); aliases=[name,*(b.get('Aliases') or [])]
        hints=[normalize_text(x) for x in (b.get('Scope Hints') or []) if normalize_text(x)]
        cur.execute('INSERT OR REPLACE INTO boxes VALUES(?,?,?,?,?,?,?,?)',(bid,normalize_text(b.get('Box Type')),name,normalize_text(b.get('Canonical Key')) or name.lower(),normalize_text(b.get('Description')),normalize_text(b.get('Status')) or 'Confirmed',json.dumps(hints,ensure_ascii=False),json.dumps(b,ensure_ascii=False,default=str)))
        for alias in dict.fromkeys(normalize_text(x) for x in aliases if normalize_text(x)):
            cur.execute('INSERT OR REPLACE INTO box_alias_map VALUES(?,?,?)',(bid,alias,alias.lower()))
        for pid in b.get('Parent Box IDs') or []:
            pid=normalize_text(pid)
            if pid: legacy_hierarchy.append((pid,bid,'legacy_parent_field'))
        for cid in b.get('Member Card IDs') or []:
            legacy_members.append((normalize_text(cid),bid,'belongs_to','Manual Configuration','','Confirmed',json.dumps({'legacy':True},ensure_ascii=False)))

    for e in box_links:
        eid=normalize_text(e.get('Box Link ID'))
        s=normalize_text(e.get('Source Box ID')); t=normalize_text(e.get('Target Box ID')); rel=normalize_text(e.get('Relation Type'))
        if not eid or not s or not t or not rel: continue
        cur.execute('INSERT OR REPLACE INTO box_edges VALUES(?,?,?,?,?,?,?,?)',(eid,s,rel,t,normalize_text(e.get('Link Basis')),normalize_text(e.get('Evidence')),normalize_text(e.get('Status')) or 'Confirmed',json.dumps(e,ensure_ascii=False,default=str)))
        if rel=='parent_of': cur.execute('INSERT OR IGNORE INTO box_hierarchy VALUES(?,?,?)',(s,t,eid))
        elif rel=='part_of': cur.execute('INSERT OR IGNORE INTO box_hierarchy VALUES(?,?,?)',(t,s,eid))

    for p,c,src in legacy_hierarchy:
        cur.execute('INSERT OR IGNORE INTO box_hierarchy VALUES(?,?,?)',(p,c,src))

    for bid in known_box_ids: cur.execute('INSERT OR IGNORE INTO box_closure VALUES(?,?,0)',(bid,bid))
    cur.execute('''INSERT OR IGNORE INTO box_closure(ancestor_id,descendant_id,depth)
                   SELECT parent_id,child_id,1 FROM box_hierarchy''')
    # Floyd-like iterative closure; min depth is kept by INSERT OR IGNORE based on first discovered path.
    changed=True
    while changed:
        before=con.total_changes
        cur.execute('''INSERT OR IGNORE INTO box_closure(ancestor_id,descendant_id,depth)
                       SELECT c.ancestor_id,h.child_id,c.depth+1
                       FROM box_closure c JOIN box_hierarchy h ON h.parent_id=c.descendant_id''')
        changed=(con.total_changes>before)

    for m in memberships:
        cid=normalize_text(m.get('Card ID')); bid=normalize_text(m.get('Box ID')); mt=normalize_text(m.get('Membership Type')) or 'belongs_to'
        if cid and bid:
            cur.execute('INSERT OR REPLACE INTO card_box VALUES(?,?,?,?,?,?,?)',(cid,bid,mt,normalize_text(m.get('Membership Basis')),normalize_text(m.get('Evidence')),normalize_text(m.get('Status')) or 'Confirmed',json.dumps(m,ensure_ascii=False,default=str)))
    for row in legacy_members: cur.execute('INSERT OR IGNORE INTO card_box VALUES(?,?,?,?,?,?,?)',row)

    for n,r in enumerate(links,1):
        lid=normalize_text(r.get('Link ID')) or f'AUTO-{n:06d}'
        cur.execute('INSERT OR REPLACE INTO links VALUES(?,?,?,?,?,?,?)',(lid,normalize_text(r.get('Source Card ID')),normalize_text(r.get('Target Card ID')),normalize_text(r.get('Link Type')),normalize_text(r.get('Link Basis')),normalize_text(r.get('Evidence')),json.dumps(r,ensure_ascii=False,default=str)))

    hashes={
        'field_xlsx':sha256_file(a.field_xlsx),'cards':sha256_file(a.cards),'evidence':sha256_file(a.evidence),'links':sha256_file(a.links),
        'context_map':sha256_file(a.context_map),'boxes':sha256_file(a.boxes),
        'box_memberships':sha256_file(a.box_memberships),'box_links':sha256_file(a.box_links)
    }
    hashes['scope_attributions']=sha256_file(a.scope_attributions)
    hashes['domains']=sha256_file(a.domains)
    meta={'schema_version':'10','tokenizer_version':TOKENIZER_VERSION,'field_count':len(fields),'card_count':len(cards),'evidence_count':len(evidence),'evidence_card_edge_count':len(evidence_links),'link_count':len(links),'business_context_edge_count':len(ctx_rows),'box_count':len(boxes),'box_membership_count':len(memberships),'box_link_count':len(box_links),'domain_count':len(domain_terms),'source_hashes':hashes}
    for k,v in meta.items(): cur.execute('INSERT INTO meta VALUES(?,?)',(k,json.dumps(v,ensure_ascii=False)))
    con.commit(); con.close()
    print(json.dumps({'output':str(a.output),**{k:v for k,v in meta.items() if k!='source_hashes'},'source_hashes':hashes},ensure_ascii=False))


def main():
    ap=argparse.ArgumentParser(description='Build AFC SQLite FTS index with Business Context, Evidence Records, Box topology/membership, Cards, Fields, and Card Links.')
    ap.add_argument('--field-xlsx'); ap.add_argument('--cards'); ap.add_argument('--evidence'); ap.add_argument('--links'); ap.add_argument('--context-map'); ap.add_argument('--scope-attributions'); ap.add_argument('--domains')
    ap.add_argument('--boxes'); ap.add_argument('--box-memberships'); ap.add_argument('--box-links'); ap.add_argument('-o','--output',required=True)
    build(ap.parse_args())

if __name__=='__main__': main()
