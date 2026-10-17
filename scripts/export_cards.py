#!/usr/bin/env python3
from pathlib import Path
import argparse, json, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from afc_card_runtime.common import read_jsonl

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('cards'); ap.add_argument('-o','--output',required=True); a=ap.parse_args()
    rows=read_jsonl(a.cards)
    headers=['Card ID','Card Type','Title','Content','Original Evidence','Language','Concept Key','Scope','Context','Box IDs','Related Card IDs','Source','Source Location','Evidence Type','Source Version','Source Date','中文補充','中文補充依據','Notes']
    wb=Workbook(); ws=wb.active; ws.title='Card_Data'; ws.append(headers)
    for c in ws[1]: c.font=Font(bold=True)
    ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:S{max(1,len(rows)+1)}'
    
    def scalar(v):
        if isinstance(v, list): return '; '.join(str(x) for x in v)
        if isinstance(v, dict): return json.dumps(v, ensure_ascii=False)
        return v
    for r in rows: ws.append([scalar(r.get(h,'未定義')) for h in headers])
    for row in ws.iter_rows():
        for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    widths=[16,14,30,60,60,12,24,20,35,25,30,35,30,16,16,16,24,22,40]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    p=Path(a.output); p.parent.mkdir(parents=True,exist_ok=True); wb.save(p); wb.close()
    wb2=load_workbook(p,read_only=True,data_only=True); ws2=wb2['Card_Data']; count=ws2.max_row-1; wb2.close()
    if count!=len(rows): raise SystemExit(f'row mismatch {count}!={len(rows)}')
    print(f'rows={count} reopen=PASS output={p}')
if __name__=='__main__': main()
