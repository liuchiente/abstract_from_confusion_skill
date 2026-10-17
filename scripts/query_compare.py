#!/usr/bin/env python3
from pathlib import Path
import argparse, json, sys
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from afc_card_runtime.common import read_jsonl, tokenize, normalize_text

FIELD_SEARCH=['名稱','中文','越文','中文補充','說明','規則','原始商品資訊','商品','商品分類','商品代號','功能代碼','System','資料來源','Related Field','來源','Supporting Source']
CARD_SEARCH=['Title','Content','Original Evidence','中文補充','Scope','Context','Card Type','Concept Key','Source','Source Location']

def load_field_xlsx(path):
    wb=load_workbook(path, read_only=True, data_only=True); ws=wb['Field_Data']
    it=ws.iter_rows(values_only=True); headers=[str(x) if x is not None else '' for x in next(it)]
    rows=[]
    for vals in it: rows.append({headers[i]:vals[i] for i in range(len(headers))})
    wb.close(); return rows

def score_record(r, query, fields, filters):
    q=normalize_text(query).lower(); toks=tokenize(query); score=0.0; reasons=[]
    for k,v in filters.items():
        if v and normalize_text(r.get(k,'')).lower()!=normalize_text(v).lower(): return None,[]
    for f in fields:
        text=normalize_text(r.get(f,'')); low=text.lower()
        if not text: continue
        if q and q in low:
            w=10 if f in {'名稱','中文','越文','中文補充','Title'} else 6; score+=w; reasons.append(f'exact:{f}')
        for t in toks:
            if t and t in low:
                w=3 if f in {'名稱','中文','越文','中文補充','Title'} else 1
                score+=w
    return score,reasons

def top(rows, query, fields, filters, k):
    scored=[]
    for r in rows:
        s,why=score_record(r,query,fields,filters)
        if s is not None and s>0: scored.append((s,r,why))
    scored.sort(key=lambda x:(-x[0], normalize_text(x[1].get('Field ID') or x[1].get('Card ID'))))
    return scored[:k]

def expand_field(seed, all_rows):
    groups={normalize_text(x[1].get('Candidate Group')) for x in seed if normalize_text(x[1].get('Candidate Group')) not in {'','未定義'}}
    ids={normalize_text(x[1].get('Field ID')) for x in seed}
    out=[]
    for r in all_rows:
        if normalize_text(r.get('Field ID')) in ids: continue
        if normalize_text(r.get('Candidate Group')) in groups: out.append(r)
    return out

def expand_cards(seed, cards, links, hops=1):
    byid={normalize_text(r.get('Card ID')):r for r in cards}; frontier={normalize_text(x[1].get('Card ID')) for x in seed}; seen=set(frontier)
    for _ in range(max(0,hops)):
        nxt=set()
        for l in links:
            s=normalize_text(l.get('Source Card ID')); t=normalize_text(l.get('Target Card ID'))
            if s in frontier and t not in seen: nxt.add(t)
            if t in frontier and s not in seen: nxt.add(s)
        seen |= nxt; frontier=nxt
        if not frontier: break
    direct={normalize_text(x[1].get('Card ID')) for x in seed}
    return [byid[i] for i in seen-direct if i in byid]

def main():
    ap=argparse.ArgumentParser(description='Deterministic candidate retrieval for AFC Field vs Card comparison.')
    ap.add_argument('--mode',choices=['field','card','compare'],default='compare')
    ap.add_argument('--query',required=True); ap.add_argument('--field-xlsx'); ap.add_argument('--cards'); ap.add_argument('--links'); ap.add_argument('--top-k',type=int,default=20); ap.add_argument('--candidate-k',type=int,default=100); ap.add_argument('--graph-hops',type=int,default=1)
    ap.add_argument('--scope'); ap.add_argument('--product'); ap.add_argument('--anchor'); ap.add_argument('-o','--output',required=True)
    a=ap.parse_args(); search_query=a.query
    if a.anchor:
        import re
        search_query=re.sub(r'(?<![A-Za-z0-9_])'+re.escape(a.anchor)+r'(?![A-Za-z0-9_])',' ',search_query,flags=re.I).strip() or a.query
    result={'query':a.query,'search_query':search_query,'anchor':a.anchor,'mode':a.mode,'top_k':a.top_k}
    if a.mode in {'field','compare'}:
        fields=load_field_xlsx(a.field_xlsx); filters={}
        if a.product or a.anchor: filters['商品代號']=a.product or a.anchor
        seed=top(fields,search_query,FIELD_SEARCH,filters,a.candidate_k); result['field']={'direct':[{'score':s,'record':r,'reasons':w} for s,r,w in seed[:a.top_k]],'expanded':expand_field(seed[:a.top_k],fields)}
    if a.mode in {'card','compare'}:
        cards=read_jsonl(a.cards); links=read_jsonl(a.links) if a.links else []; filters={}
        if a.scope or a.anchor: filters['Scope']=a.scope or a.anchor
        seed=top(cards,search_query,CARD_SEARCH,filters,a.candidate_k); result['card']={'direct':[{'score':s,'record':r,'reasons':w} for s,r,w in seed[:a.top_k]],'expanded':expand_cards(seed[:a.top_k],cards,links,a.graph_hops)}
    Path(a.output).parent.mkdir(parents=True,exist_ok=True); Path(a.output).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8'); print(a.output)
if __name__=='__main__': main()
