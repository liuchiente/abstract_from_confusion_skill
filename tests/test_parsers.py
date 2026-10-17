"""Source-structure ingestion tests.

v0.4.1 incremented its `discovered` and `extracted` counters on the same line,
so SOURCE_PARSING_GAP could only ever fire on blank paragraphs while being
structurally unable to detect a missed table, text box or shape -- and reported
PASS for content no parser reads at all (header/footer text, footnotes,
comments, charts).
"""
import json
import struct
import subprocess
import sys
import zipfile
import zlib
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def run(args, expect_ok=True):
    proc = subprocess.run([sys.executable, *args], capture_output=True, text=True, cwd=ROOT)
    if expect_ok:
        assert proc.returncode == 0, proc.stdout + proc.stderr
    return proc


def tiny_png():
    def chunk(tag, data):
        body = tag + data
        return struct.pack('>I', len(data)) + body + struct.pack('>I', zlib.crc32(body) & 0xFFFFFFFF)
    return (b'\x89PNG\r\n\x1a\n'
            + chunk(b'IHDR', struct.pack('>IIBBBBB', 1, 1, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(b'\x00\xff\x00\x00'))
            + chunk(b'IEND', b''))


@pytest.fixture
def docx_with_everything(tmp_path):
    docx = pytest.importorskip('docx')
    png = tmp_path / 'p.png'
    png.write_bytes(tiny_png())
    d = docx.Document()
    d.add_heading('CA 新增投保流程', 1)
    for _ in range(4):
        d.add_paragraph('')                      # blank layout paragraphs
    d.add_paragraph('確認頁按下返回時保留已輸入資料。')
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = 'Error Message'
    t.cell(1, 0).text = '加買失敗'
    outer = d.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=2)   # nested table, parent row empty
    inner.cell(0, 0).text = 'Tải lên thất bại'
    d.add_paragraph('').add_run().add_picture(str(png))
    d.sections[0].header.paragraphs[0].text = '機密 — CA 規格書 v2'
    src = tmp_path / 'in'
    src.mkdir()
    path = src / 'spec.docx'
    d.save(path)
    # Inject footnote and comment parts, as a real reviewed document would have.
    tmp = tmp_path / 'tmp.docx'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr('word/footnotes.xml',
                      f'<?xml version="1.0"?><w:footnotes xmlns:w="{W}"><w:footnote w:id="1">'
                      f'<w:p><w:r><w:t>註腳：本規則自 2024 年起適用</w:t></w:r></w:p></w:footnote></w:footnotes>')
        zout.writestr('word/comments.xml',
                      f'<?xml version="1.0"?><w:comments xmlns:w="{W}"><w:comment w:id="1" w:author="PM">'
                      f'<w:p><w:r><w:t>這條規則需要確認</w:t></w:r></w:p></w:comment></w:comments>')
    path.unlink()
    tmp.rename(path)
    return src


def parse_docx(tmp_path, src):
    units = tmp_path / 'units.jsonl'
    report = tmp_path / 'structure.json'
    run(['scripts/parse_word.py', str(src), '-o', str(units),
         '--structure-report', str(report), '--media-dir', str(tmp_path / 'media')])
    return units, json.loads(report.read_text(encoding='utf-8'))


def test_blank_paragraphs_are_not_a_parsing_gap(tmp_path, docx_with_everything):
    _, report = parse_docx(tmp_path, docx_with_everything)
    assert report['parser_gaps'] == []
    assert not any(g['structure'] == 'body_paragraphs_with_text' for g in report['parser_gaps'])


def test_discovery_is_independent_of_extraction(tmp_path, docx_with_everything):
    _, report = parse_docx(tmp_path, docx_with_everything)
    f = report['files'][0]
    # Counts come from two separate passes, so they can disagree -- and here they
    # agree because nothing was actually missed.
    assert f['discovered']['tables'] == f['extracted']['tables'] == 3
    assert f['discovered']['nested_tables'] == f['extracted']['nested_tables'] == 1
    assert f['discovered_image_parts'] and f['extracted_image_parts']


def test_unreadable_structures_are_named_not_passed(tmp_path, docx_with_everything):
    _, report = parse_docx(tmp_path, docx_with_everything)
    kinds = {u['structure'] for u in report['unsupported_structures']}
    assert {'header_footer_text', 'footnote_text', 'comment_text'} <= kinds
    assert report['status'] == 'REVIEW'


def test_audit_routes_unsupported_separately_from_a_gap(tmp_path, docx_with_everything):
    units, _ = parse_docx(tmp_path, docx_with_everything)
    proc = run(['scripts/audit_source_structure.py', '--source-map', str(units),
                '--structure-report', str(tmp_path / 'structure.json'),
                '--expect-parsers', 'parse_word.py'])
    audit = json.loads(proc.stdout)
    assert 'SOURCE_PARSING_GAP' not in audit['recommended_routes']
    assert 'SOURCE_UNSUPPORTED_STRUCTURE' in audit['recommended_routes']
    assert audit['status'] == 'REVIEW'
    assert audit['parser_gaps'] == []


def test_a_missing_parser_report_is_a_gap(tmp_path, docx_with_everything):
    units, _ = parse_docx(tmp_path, docx_with_everything)
    proc = run(['scripts/audit_source_structure.py', '--source-map', str(units),
                '--structure-report', str(tmp_path / 'structure.json'),
                '--expect-parsers', 'parse_word.py,parse_excel.py'], expect_ok=False)
    audit = json.loads(proc.stdout)
    assert audit['status'] == 'FAIL'
    assert 'parse_excel.py' in audit['missing_parser_reports']


def test_nested_table_rows_survive_with_structure_paths(tmp_path, docx_with_everything):
    units, _ = parse_docx(tmp_path, docx_with_everything)
    rows = [json.loads(l) for l in units.read_text(encoding='utf-8').splitlines() if l.strip()]
    nested = [r for r in rows if r.get('is_nested_table')]
    assert nested, 'nested table rows must become Source Units'
    assert 'Tải lên thất bại' in nested[0]['value']
    assert 'Nested 1' in nested[0]['structure_path_text']


@pytest.fixture
def xlsx_workbook(tmp_path):
    openpyxl = pytest.importorskip('openpyxl')
    import datetime
    from openpyxl import Workbook
    src = tmp_path / 'in'
    (src / 'sub').mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Spec'
    ws['A1'], ws['B1'], ws['C1'] = '欄位名稱', 'Error Message', 'Component Type'
    ws.merge_cells('A2:A4')
    ws['A2'] = '投保資料'
    ws['B2'], ws['C2'] = 'Tải lên thất bại', 'textbox'
    ws['B3'], ws['C3'] = '加買失敗', 'textbox'
    ws['B4'], ws['C4'] = '無法加買', 'dropdown'
    ws['A6'] = '生效日'
    ws['B6'] = datetime.datetime(2024, 1, 1)
    ws['B6'].number_format = 'yyyy-mm-dd'
    ws['A7'] = '合併值'
    ws['B7'] = '=A1&"-"&B1'
    hidden = wb.create_sheet('Hidden')
    hidden['A1'] = '隱藏內容'
    hidden.sheet_state = 'hidden'
    wb.save(src / 'spec.xlsx')
    wb2 = Workbook()
    wb2.active['A1'] = '另一份同名檔案'
    wb2.save(src / 'sub' / 'spec.xlsx')
    return src


def parse_xlsx(tmp_path, src, extra=()):
    units = tmp_path / 'xunits.jsonl'
    report = tmp_path / 'xstructure.json'
    run(['scripts/parse_excel.py', str(src), '-o', str(units),
         '--structure-report', str(report), '--media-dir', str(tmp_path / 'xmedia'), *extra])
    rows = [json.loads(l) for l in units.read_text(encoding='utf-8').splitlines() if l.strip()]
    return rows, json.loads(report.read_text(encoding='utf-8'))


def test_merged_header_reaches_the_rows_it_governs(tmp_path, xlsx_workbook):
    rows, _ = parse_xlsx(tmp_path, xlsx_workbook)
    merged = [r for r in rows if r.get('has_merged_inheritance')]
    assert len(merged) == 2
    assert all('投保資料' in r['value'] for r in merged)


def test_dates_are_rendered_as_displayed(tmp_path, xlsx_workbook):
    rows, _ = parse_xlsx(tmp_path, xlsx_workbook)
    row6 = next(r for r in rows if r.get('row') == 6)
    assert '2024-01-01' in row6['value']
    assert '00:00:00' not in row6['value']


def test_a_formula_without_a_cached_value_is_kept_and_flagged(tmp_path, xlsx_workbook):
    rows, report = parse_xlsx(tmp_path, xlsx_workbook)
    row7 = next(r for r in rows if r.get('row') == 7)
    assert row7['has_uncached_formula']
    cell = next(c for c in row7['cells'] if c.get('formula'))
    assert cell['value_status'] == 'NO_CACHED_VALUE'
    assert any(u['structure'] == 'formula_without_cached_value' for u in report['unsupported_structures'])


def test_same_named_workbooks_do_not_collide(tmp_path, xlsx_workbook):
    rows, _ = parse_xlsx(tmp_path, xlsx_workbook)
    ids = [r['source_unit_id'] for r in rows]
    assert len(ids) == len(set(ids))
    assert any(i.startswith('sub/spec.xlsx::') for i in ids)
    assert any(i.startswith('spec.xlsx::') for i in ids)


def test_hidden_sheets_are_reported_not_silently_dropped(tmp_path, xlsx_workbook):
    rows, report = parse_xlsx(tmp_path, xlsx_workbook)
    assert not any(r.get('sheet') == 'Hidden' for r in rows)
    assert any(u['structure'] == 'skipped_hidden_sheet' for u in report['unsupported_structures'])
    rows2, _ = parse_xlsx(tmp_path, xlsx_workbook, extra=['--include-hidden-sheets'])
    hidden_rows = [r for r in rows2 if r.get('sheet') == 'Hidden']
    assert hidden_rows and hidden_rows[0]['sheet_hidden'] is True
